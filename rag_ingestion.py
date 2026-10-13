from pathlib import Path
from sentence_transformers import SentenceTransformer
import fitz
from docx import Document
import csv

def read_pdf(file_path):
    text = ""

    with fitz.open(file_path) as pdf:
        for page in pdf:
            text += page.get_text()

    return text

def read_txt(file_path):
    with open(file_path, "r", encoding="utf-8") as file:
        return file.read()

def read_docx(file_path):
    document = Document(file_path)
    text = ""

    for paragraph in document.paragraphs:
        text += paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            row_text = []

        for cell in row.cells:
            row_text.append(cell.text)

        text += " |".join(row_text) + "\n"

    return text

def read_csv(file_path):
    text = ""

    with open(file_path, "r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)

        for row in reader:
            row_text = " | ".join(
                f"{column}: {value}"
                for column, value in row.items()
            )
            text += row_text + "\n"

    return text

LOADERS = {
    ".pdf": read_pdf,
    ".txt": read_txt,
    ".docx" : read_docx,
    ".csv": read_csv
}

def load_documents():
    documents = []

    for file_path in Path("knowledge_base").rglob("*"):
        if not file_path.is_file():
            continue

        extension = file_path.suffix.lower()
        reader = LOADERS.get(extension)

        if reader is None:
            continue

        text = reader(file_path)

        if text.strip():
            documents.append({
                "source": str(file_path),
                "text": text,
                "category": file_path.parent.name
            })

    return documents


def chunk_text(text):
    paragraphs = text.split("\n")

    chunks = []

    for paragraph in paragraphs:
        if paragraph.strip():
            chunks.append(paragraph.strip())

    return chunks

def generate_embeddings(text_chunks):
    embeddings = model.encode(
        text_chunks,
        batch_size=32,
        show_progress_bar = True
    )
    return embeddings

model = SentenceTransformer("all-MiniLM-L6-v2")



